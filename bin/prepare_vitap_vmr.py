#!/usr/bin/env python3
"""Convert an ICTV VMR workbook or CSV into VITAP's nine-column input."""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
from pathlib import Path
from typing import Iterable, Mapping


OUTPUT_COLUMNS = (
    "Virus GENBANK accession",
    "Realm",
    "Kingdom",
    "Phylum",
    "Class",
    "Order",
    "Family",
    "Genus",
    "Species",
)


def normalize_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def column_mapping(headers: Iterable[object]) -> dict[str, int]:
    normalized = [normalize_header(header) for header in headers]
    mapping: dict[str, int] = {}
    for output_column in OUTPUT_COLUMNS:
        wanted = normalize_header(output_column)
        matches = [i for i, value in enumerate(normalized) if value == wanted]
        if output_column == "Virus GENBANK accession" and not matches:
            matches = [
                i
                for i, value in enumerate(normalized)
                if value.startswith("virus genbank accession")
            ]
        if len(matches) != 1:
            raise ValueError(
                f"Expected exactly one '{output_column}' column; found {len(matches)}."
            )
        mapping[output_column] = matches[0]
    return mapping


def read_csv_rows(path: Path) -> tuple[str, list[dict[str, str]]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        try:
            headers = next(reader)
        except StopIteration as exc:
            raise ValueError("The supplied VMR CSV is empty.") from exc
        mapping = column_mapping(headers)
        rows = []
        for values in reader:
            rows.append(
                {
                    column: str(values[index]).strip() if index < len(values) else ""
                    for column, index in mapping.items()
                }
            )
    return "csv", rows


def read_xlsx_rows(path: Path, requested_sheet: str | None) -> tuple[str, list[dict[str, str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read ICTV Excel workbooks.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheets = (
        [workbook[requested_sheet]] if requested_sheet else list(workbook.worksheets)
    )
    failures: list[str] = []
    for worksheet in worksheets:
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=1, max_row=30, values_only=True), start=1
        ):
            try:
                mapping = column_mapping(values)
            except ValueError as exc:
                failures.append(f"{worksheet.title} row {row_number}: {exc}")
                continue

            rows: list[dict[str, str]] = []
            for data in worksheet.iter_rows(min_row=row_number + 1, values_only=True):
                rows.append(
                    {
                        column: str(data[index]).strip()
                        if index < len(data) and data[index] is not None
                        else ""
                        for column, index in mapping.items()
                    }
                )
            workbook.close()
            return worksheet.title, rows

    workbook.close()
    detail = failures[-1] if failures else "no worksheets were available"
    raise ValueError(f"Could not locate the nine VITAP columns ({detail}).")


def inspect_accessions(value: str) -> tuple[int, int, int, list[str]]:
    entry_count = coordinate_count = labeled_count = 0
    invalid: list[str] = []
    for raw_entry in value.split(";"):
        entry = raw_entry.strip()
        if not entry:
            continue
        entry_count += 1
        if ":" in entry:
            labeled_count += 1
            entry = entry.split(":", 1)[1].strip()
        coordinate_match = re.fullmatch(r"(.+?)\s*\((\d+)\.(\d+)\)", entry)
        if coordinate_match:
            coordinate_count += 1
            entry = coordinate_match.group(1).strip()
        # Match the practical GenBank/RefSeq syntax without over-constraining
        # future accession prefixes. VITAP removes version suffixes itself.
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*(?:\.\d+)?", entry):
            invalid.append(raw_entry.strip())
    return entry_count, coordinate_count, labeled_count, invalid


def infer_label(path: Path, explicit_label: str | None) -> str:
    if explicit_label:
        label = explicit_label.strip()
    else:
        match = re.search(r"MSL[_.-]?(\d+)", path.name, flags=re.IGNORECASE)
        if not match:
            raise ValueError(
                "Could not infer an MSL release from the VMR filename; "
                "supply --label (for example, VMR-MSL41)."
            )
        label = f"VMR-MSL{match.group(1)}"
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", label):
        raise ValueError(f"Invalid database label '{label}'.")
    return label


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_metadata(path: Path, values: Mapping[str, object]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(values), delimiter="\t")
        writer.writeheader()
        writer.writerow(values)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-csv", required=True, type=Path)
    parser.add_argument("--metadata", required=True, type=Path)
    parser.add_argument("--source-url", default="user-supplied")
    parser.add_argument("--sheet")
    parser.add_argument("--label")
    args = parser.parse_args()

    source = args.input.resolve()
    if not source.is_file():
        raise FileNotFoundError(f"VMR source not found: {source}")
    label = infer_label(source, args.label)
    if source.suffix.lower() == ".csv":
        worksheet, rows = read_csv_rows(source)
    elif source.suffix.lower() in {".xlsx", ".xlsm"}:
        worksheet, rows = read_xlsx_rows(source, args.sheet)
    else:
        raise ValueError("The VMR source must be .xlsx, .xlsm, or .csv.")

    written: list[dict[str, str]] = []
    accession_entries = coordinate_entries = labeled_entries = 0
    invalid_entries: list[str] = []
    for row in rows:
        accession = row["Virus GENBANK accession"].strip()
        if not accession:
            continue
        entries, coordinates, labels, invalid = inspect_accessions(accession)
        accession_entries += entries
        coordinate_entries += coordinates
        labeled_entries += labels
        invalid_entries.extend(invalid)
        written.append({column: row[column].strip() for column in OUTPUT_COLUMNS})

    if invalid_entries:
        examples = ", ".join(repr(value) for value in invalid_entries[:10])
        raise ValueError(
            f"Found {len(invalid_entries)} malformed accession entries; examples: {examples}"
        )
    if not written:
        raise ValueError("The VMR contains no rows with GenBank accessions.")

    args.output_csv.parent.mkdir(parents=True, exist_ok=True)
    with args.output_csv.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(written)

    write_metadata(
        args.metadata,
        {
            "release_label": label,
            "source_file": source.name,
            "source_url": args.source_url,
            "source_sha256": sha256(source),
            "worksheet": worksheet,
            "input_rows": len(rows),
            "written_rows": len(written),
            "skipped_rows_without_accession": len(rows) - len(written),
            "accession_entries": accession_entries,
            "labeled_entries": labeled_entries,
            "coordinate_entries": coordinate_entries,
            "validation": "passed",
        },
    )


if __name__ == "__main__":
    main()
